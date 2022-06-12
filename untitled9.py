# -*- coding: utf-8 -*-
"""
Created on Sun May 15 17:14:57 2022

@author: edgar
"""

from openpyxl import Workbook, load_workbook
from openpyxl.chart import PieChart, Reference
class Aluno:
    def __init__(self, nome="", nota_1=0.0, nota_2=0.0, nota_3=0.0):
        self.nome = nome
        self.nota_1 = nota_1
        self.nota_2 = nota_2
        self.nota_3 = nota_3
        self.media = self.calcula_media()
        self.status = self.verifica_aluno()
    def calcula_media(self):
        notas=[self.nota_1,self.nota_2,self.nota_3]
        x= max(notas)
        notas.remove(x)
        result = (3*x+2*notas[0]+2*notas[1])/7
        return result
    
    def verifica_aluno(self):
        if self.media >6:
            self.status = "Aprovado"
        else:
            self.status = "Reprovado"
        return self.status

    def to_list(self):
        return(self.nome,self.nota_1,self.nota_2,self.nota_3,self.media,self.status)


class Alunos:
    def __init__(self,lista_alunos=[]):
        self.alunos = lista_alunos.copy()

    def insere(self,aluno):
        self.alunos.append(aluno)

    def __str__(self):
        txt = "["
        for aluno in self.alunos:
            txt = txt + aluno.nome + "," + str(aluno.nota_1) + "," + str(aluno.nota_2) + "," + str(aluno.nota_3) + "," + str(aluno.media) + "," + aluno.status + ","
            txt = txt + "]"
        return txt

    def salva_excel(self, nom_workbook="Alunos.xlsx",
        nom_planilha="planilha1",cabecalho=[],celula_inicial=""):
        arquivo_excel = Workbook()
        planilha1 = arquivo_excel.active
        planilha1.title = nom_planilha
        
        planilha1.append(cabecalho)
        Aprovados=[]
        Reprovados=[]
        for aluno in self.alunos:
            if aluno.status =="Aprovado":
                Aprovados.append(aluno)
            else:
                Reprovados.append(aluno)
                planilha1.append(aluno.to_list())
                A= len(Aprovados)
                R= len(Reprovados)
                planilha1['H1']="Aprovados"
                planilha1['H2']=A
                planilha1['G1']="Reprovados"
                planilha1['G2']=R
                chart = PieChart()
                labels = Reference(planilha1, min_row = 1,
                min_col = 7, max_col = 8)
                data = Reference(planilha1, min_row = 2,
                min_col = 7, max_col = 8)
                chart.add_data(data, titles_from_data = True)
                chart.set_categories(labels)
                chart.title = " Aprovados/Reprovados "
                planilha1.add_chart(chart, "J2")
                arquivo_excel.save(nom_workbook)
        return "Planilha criada"

cabecalho = ["Nome","Nota 1","Nota 2","Nota 3","Média", "Situação"]
aluno_1 = Aluno("Pedro Lucas Santos", 8.0, 6.0, 10.0)
aluno_2 = Aluno("José Luis Silva", 3.0, 5.0, 4.0)
aluno_3 = Aluno("João Pedro Almeida", 8.0, 9.0, 4.0)
alunos = Alunos()
alunos.insere(aluno_1)
alunos.insere(aluno_2)
alunos.insere(aluno_3)
alunos.salva_excel("Alunos_teste.xlsx","Alunos", cabecalho,"A1")
print(alunos)